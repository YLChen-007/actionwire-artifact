#!/usr/bin/env python3
"""Generate per-project vulnerability, received, and CVE/GHSA totals."""

from __future__ import annotations

import argparse
import os
import re
import stat
import sys
import tempfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Sequence

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_WORKBOOK = (
    REPO_ROOT
    / "design/study/vuls-new-found/all-new-xclaw-vulnerabilities.xlsx"
)
DEFAULT_OUTPUT = REPO_ROOT / "design/paperdata/cve-data.md"
DEFAULT_SHEET = "all-new-xclaw-vulnerabilities"
PAPER_PROJECT_ORDER = (
    "hermes-agent",
    "nanobot",
    "chatgpt-on-wechat",
    "AstrBot",
    "QwenPaw",
    "poco-agent",
    "openclaw",
    "nanoclaw",
    "openclaw-cn",
    "mercury-agent",
    "droidclaw",
    "lettabot",
)
REQUIRED_COLUMNS = {"project", "received", "cve-num", "report-path"}
ASSIGNED_ID_PATTERN = re.compile(
    r"\b(?:CVE-\d{4}-\d{4,}|GHSA-[0-9A-Z]{4}-[0-9A-Z]{4}-[0-9A-Z]{4})\b",
    re.IGNORECASE,
)


class CveDataError(ValueError):
    """Raised when the workbook cannot produce trustworthy CVE totals."""


@dataclass
class ProjectCounts:
    vulnerabilities: int = 0
    received: int = 0
    assigned_ids: set[str] = field(default_factory=set)


@dataclass(frozen=True)
class CveSummary:
    projects: tuple[tuple[str, int, int, int], ...]
    vulnerability_rows: int
    received_rows: int
    assigned_cve_ids: int
    distinct_report_paths: int
    duplicate_report_paths: int
    duplicate_rows: int
    blank_report_paths: int


def _text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _display_path(path: Path) -> str:
    resolved = path.resolve()
    try:
        return str(resolved.relative_to(REPO_ROOT))
    except ValueError:
        return str(resolved)


def collect_cve_summary(workbook_path: Path, sheet_name: str) -> CveSummary:
    if not workbook_path.is_file():
        raise CveDataError(f"workbook does not exist: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise CveDataError(
                f"missing sheet {sheet_name!r}; available sheets: {workbook.sheetnames}"
            )
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows()
        header = next(rows, None)
        if header is None:
            raise CveDataError(f"sheet {sheet_name!r} is empty")

        names = [_text(cell.value) for cell in header]
        populated_names = [name for name in names if name]
        if len(populated_names) != len(set(populated_names)):
            duplicates = sorted(
                name for name, count in Counter(populated_names).items() if count > 1
            )
            raise CveDataError(f"duplicate header columns: {duplicates}")
        columns = {name: index for index, name in enumerate(names) if name}
        missing = sorted(REQUIRED_COLUMNS - set(columns))
        if missing:
            raise CveDataError(f"missing required columns: {missing}")

        projects: dict[str, ProjectCounts] = {}
        report_paths: Counter[str] = Counter()
        blank_report_paths = 0
        vulnerability_rows = 0

        for row_number, row in enumerate(rows, start=2):
            values = [cell.value for cell in row]
            project = _text(values[columns["project"]])
            if not project:
                if any(value is not None and _text(value) for value in values):
                    raise CveDataError(
                        f"nonempty workbook row {row_number} has no project"
                    )
                continue

            counts = projects.setdefault(project, ProjectCounts())
            counts.vulnerabilities += 1
            vulnerability_rows += 1

            if _text(values[columns["received"]]):
                counts.received += 1

            report_path = _text(values[columns["report-path"]]).replace("\\", "/")
            if report_path:
                report_paths[report_path] += 1
            else:
                blank_report_paths += 1

            cve_value = _text(values[columns["cve-num"]])
            if cve_value:
                assigned_ids = {
                    match.upper() for match in ASSIGNED_ID_PATTERN.findall(cve_value)
                }
                counts.assigned_ids.update(assigned_ids)
    finally:
        workbook.close()

    duplicate_report_paths = sum(count > 1 for count in report_paths.values())
    duplicate_rows = sum(count - 1 for count in report_paths.values() if count > 1)
    project_rows = []
    for project in PAPER_PROJECT_ORDER:
        counts = projects.get(project, ProjectCounts())
        project_rows.append(
            (
                project,
                counts.vulnerabilities,
                counts.received,
                len(counts.assigned_ids),
            )
        )
    paper_projects = set(PAPER_PROJECT_ORDER)
    project_rows.extend(
        (project, counts.vulnerabilities, counts.received, len(counts.assigned_ids))
        for project, counts in projects.items()
        if project not in paper_projects
    )
    ordered_project_rows = tuple(project_rows)
    return CveSummary(
        projects=ordered_project_rows,
        vulnerability_rows=vulnerability_rows,
        received_rows=sum(row[2] for row in ordered_project_rows),
        assigned_cve_ids=sum(row[3] for row in ordered_project_rows),
        distinct_report_paths=len(report_paths),
        duplicate_report_paths=duplicate_report_paths,
        duplicate_rows=duplicate_rows,
        blank_report_paths=blank_report_paths,
    )


def render_markdown(
    summary: CveSummary, workbook_path: Path, sheet_name: str
) -> str:
    lines = [
        "# Vulnerabilities and assigned CVE/GHSA IDs",
        "",
        f"Source: `{_display_path(workbook_path)}`, sheet `{sheet_name}`.",
        "",
        "Regenerate from the repository root:",
        "",
        "```bash",
        "python design/paperdata/script/generate_cve_data.py",
        "```",
        "",
        "Each populated project row is counted as one vulnerability. `Received` counts "
        "rows with a nonblank `received` cell. Assigned CVE/GHSA IDs are distinct "
        "`CVE-YYYY-NNNN...` or `GHSA-xxxx-xxxx-xxxx` identifiers in `cve-num` for "
        "that project.",
        "All paper projects are emitted in table order; projects absent from the "
        "sheet receive zero values, and sheet-only projects are appended.",
        "",
        "| Project | Vulnerabilities | Received | Assigned CVE/GHSA IDs |",
        "|---|---:|---:|---:|",
    ]
    lines.extend(
        f"| {project} | {vulnerabilities} | {received} | {assigned_cves} |"
        for project, vulnerabilities, received, assigned_cves in summary.projects
    )
    lines.extend(
        [
            f"| **Total** | **{summary.vulnerability_rows}** | "
            f"**{summary.received_rows}** | "
            f"**{summary.assigned_cve_ids}** |",
            "",
            f"Duplicate audit: **{summary.duplicate_report_paths}** report paths are "
            f"repeated, accounting for **{summary.duplicate_rows}** additional rows. "
            f"Deduplicating nonblank `report-path` values gives "
            f"**{summary.distinct_report_paths}** unique report paths.",
        ]
    )
    if summary.blank_report_paths:
        lines.extend(
            [
                "",
                f"Rows with a blank `report-path`: **{summary.blank_report_paths}**.",
            ]
        )
    return "\n".join(lines) + "\n"


def write_output(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    mode = stat.S_IMODE(path.stat().st_mode) if path.exists() else 0o644
    descriptor, temporary_name = tempfile.mkstemp(
        dir=path.parent, prefix=f".{path.name}.", suffix=".tmp"
    )
    temporary_path = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as handle:
            handle.write(content)
        temporary_path.chmod(mode)
        temporary_path.replace(path)
    except BaseException:
        temporary_path.unlink(missing_ok=True)
        raise


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate per-project vulnerability and assigned-CVE/GHSA totals."
    )
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        summary = collect_cve_summary(args.workbook, args.sheet)
        write_output(
            args.output,
            render_markdown(summary, args.workbook, args.sheet),
        )
    except (CveDataError, OSError, ValueError) as error:
        print(f"error: {error}", file=sys.stderr)
        return 1

    print(
        f"wrote {args.output}: {summary.vulnerability_rows} vulnerabilities, "
        f"{summary.received_rows} received, "
        f"{summary.assigned_cve_ids} assigned CVE/GHSA IDs"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
