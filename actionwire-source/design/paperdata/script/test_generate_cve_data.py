#!/usr/bin/env python3
"""Tests for the CVE-data Markdown generator."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

import generate_cve_data as cve_data  # noqa: E402


class CveDataTests(unittest.TestCase):
    def create_workbook(self, directory: str) -> Path:
        path = Path(directory) / "fixture.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = cve_data.DEFAULT_SHEET
        sheet.append(["project", "received", "cve-num", "report-path"])
        sheet.append(["project-a", "accepted", "CVE-2026-1000", "/reports/a.md"])
        sheet.append(["project-a", None, None, "/reports/a.md"])
        sheet.append(
            [
                "project-b",
                "2026-01-01: received",
                "CVE-2026-2000; CVE-2026-2001; CVE-2026-2001",
                "/reports/b.md",
            ]
        )
        sheet.append([None, None, None, None])
        workbook.save(path)
        workbook.close()
        return path

    def test_counts_projects_cves_and_duplicate_paths(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            summary = cve_data.collect_cve_summary(
                self.create_workbook(directory), cve_data.DEFAULT_SHEET
            )
        rows = {row[0]: row[1:] for row in summary.projects}
        self.assertEqual((2, 1, 1), rows["project-a"])
        self.assertEqual((1, 1, 2), rows["project-b"])
        self.assertEqual((0, 0, 0), rows["poco-agent"])
        self.assertEqual((0, 0, 0), rows["lettabot"])
        self.assertEqual(3, summary.vulnerability_rows)
        self.assertEqual(2, summary.received_rows)
        self.assertEqual(3, summary.assigned_cve_ids)
        self.assertEqual(2, summary.distinct_report_paths)
        self.assertEqual(1, summary.duplicate_report_paths)
        self.assertEqual(1, summary.duplicate_rows)

    def test_rendered_table_has_totals_and_reproduction_command(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook = self.create_workbook(directory)
            summary = cve_data.collect_cve_summary(
                workbook, cve_data.DEFAULT_SHEET
            )
            rendered = cve_data.render_markdown(
                summary, workbook, cve_data.DEFAULT_SHEET
            )
        self.assertIn("| project-a | 2 | 1 | 1 |", rendered)
        self.assertIn("| **Total** | **3** | **2** | **3** |", rendered)
        self.assertIn("Assigned CVE/GHSA IDs", rendered)
        self.assertNotIn("Assigned CVE IDs", rendered)
        self.assertIn(
            "python design/paperdata/script/generate_cve_data.py", rendered
        )

    def test_projects_follow_selected_paper_order(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "order.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = cve_data.DEFAULT_SHEET
            sheet.append(["project", "received", "cve-num", "report-path"])
            for project in ("openclaw", "LobsterAI", "nanobot", "hermes-agent"):
                sheet.append([project, None, None, f"/{project}.md"])
            workbook.save(path)
            workbook.close()
            summary = cve_data.collect_cve_summary(path, cve_data.DEFAULT_SHEET)
        self.assertEqual(
            [*cve_data.PAPER_PROJECT_ORDER, "LobsterAI"],
            [row[0] for row in summary.projects],
        )

    def test_ghsa_identifier_is_counted_as_assigned_id(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = self.create_workbook(directory)
            workbook = load_workbook(path)
            workbook[cve_data.DEFAULT_SHEET]["C2"] = "GHSA-xxxx-yyyy-zzzz"
            workbook.save(path)
            workbook.close()
            summary = cve_data.collect_cve_summary(path, cve_data.DEFAULT_SHEET)
        rows = {row[0]: row[1:] for row in summary.projects}
        self.assertEqual((2, 1, 1), rows["project-a"])
        self.assertEqual((1, 1, 2), rows["project-b"])

    def test_required_columns_are_validated(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bad.xlsx"
            workbook = Workbook()
            workbook.active.append(["project", "received", "cve-num"])
            workbook.save(path)
            workbook.close()
            with self.assertRaisesRegex(cve_data.CveDataError, "report-path"):
                cve_data.collect_cve_summary(path, "Sheet")


if __name__ == "__main__":
    unittest.main()
